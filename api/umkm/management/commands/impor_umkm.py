"""
Importer dataset UMKM dari file CSV atau XLSX.

Contoh pakai:
    python manage.py impor_umkm datasets/pontianak.xlsx --kota Pontianak --provinsi "Kalimantan Barat"
    python manage.py impor_umkm data.csv --dry-run
    python manage.py impor_umkm data.csv --publish

Data hasil impor masuk sebagai `draft` (kecuali pakai --publish),
supaya bisa diverifikasi dulu lewat panel admin.
"""

import csv
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from umkm.models import Kategori, Umkm

# Nama kolom yang mungkin dipakai di dataset -> field kita
ALIASES = {
    "nama": ["nama usaha", "nama_usaha", "nama umkm", "nama pemohon", "nama", "usaha"],
    "kategori": ["kategori usaha", "kegiatan usaha", "jenis usaha", "bidang usaha", "kategori"],
    "alamat": ["alamat usaha", "alamat"],
    "kelurahan": ["kelurahan usaha", "kelurahan", "desa", "gampong"],
    "kecamatan": ["kecamatan usaha", "kecamatan"],
    "kota": ["kota", "kabupaten", "kota/kabupaten"],
    "provinsi": ["provinsi"],
    "deskripsi": ["deskripsi", "keterangan"],
}


def normalisasi(teks: str) -> str:
    return " ".join(str(teks or "").strip().lower().split())


def baca_baris(path: Path):
    """Balikin (headers, list-of-dict). Mendukung .csv dan .xlsx."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return [], []
        headers = [normalisasi(h) for h in rows[0]]
        data = []
        for r in rows[1:]:
            data.append({headers[i]: r[i] for i in range(len(headers))})
        return headers, data

    # default: CSV
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = [normalisasi(h) for h in (reader.fieldnames or [])]
        data = []
        for row in reader:
            data.append({normalisasi(k): v for k, v in row.items()})
    return headers, data


def peta_kolom(headers: list[str]) -> dict[str, str]:
    """Cocokkan field kita dengan nama kolom asli yang ada di file."""
    peta: dict[str, str] = {}
    for field, alias_list in ALIASES.items():
        for alias in alias_list:
            if alias in headers:
                peta[field] = alias
                break
    return peta


class Command(BaseCommand):
    help = "Impor data UMKM dari file CSV/XLSX ke database."

    def add_arguments(self, parser):
        parser.add_argument("file", type=str, help="Path ke file CSV atau XLSX")
        parser.add_argument("--kota", default="", help="Isi kolom kota kalau tidak ada di file")
        parser.add_argument("--provinsi", default="", help="Isi kolom provinsi kalau tidak ada di file")
        parser.add_argument("--sumber", default="", help="Nama sumber data, mis. dataset:pontianak")
        parser.add_argument("--publish", action="store_true", help="Langsung tandai terverifikasi")
        parser.add_argument("--dry-run", action="store_true", help="Cuma tampilkan, tidak menyimpan")

    def handle(self, *args, **options):
        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"File tidak ketemu: {path}")

        headers, rows = baca_baris(path)
        if not headers:
            raise CommandError("File kosong atau header tidak terbaca.")

        peta = peta_kolom(headers)
        if "nama" not in peta:
            raise CommandError(
                "Kolom nama usaha tidak ketemu. "
                f"Header yang terbaca: {', '.join(headers)}"
            )

        self.stdout.write(f"Kolom terdeteksi: {peta}")

        dibuat = 0
        dilewati = 0
        dry = options["dry_run"]
        status = Umkm.Status.TERVERIFIKASI if options["publish"] else Umkm.Status.DRAFT

        for baris in rows:
            def ambil(field):
                kolom = peta.get(field)
                nilai = baris.get(kolom) if kolom else None
                return str(nilai).strip() if nilai is not None else ""

            nama = ambil("nama")
            if not nama:
                dilewati += 1
                continue

            kategori_nama = ambil("kategori")
            kota = ambil("kota") or options["kota"]
            provinsi = ambil("provinsi") or options["provinsi"]

            if dry:
                self.stdout.write(f"[dry-run] {nama} | {kategori_nama} | {kota}")
                dibuat += 1
                continue

            kategori = None
            if kategori_nama:
                kategori, _ = Kategori.objects.get_or_create(nama=kategori_nama)

            Umkm.objects.update_or_create(
                nama=nama,
                kota=kota,
                defaults={
                    "kategori": kategori,
                    "deskripsi": ambil("deskripsi"),
                    "alamat": ambil("alamat"),
                    "kelurahan": ambil("kelurahan"),
                    "kecamatan": ambil("kecamatan"),
                    "provinsi": provinsi,
                    "status": status,
                    "sumber": options["sumber"] or f"impor:{path.name}",
                },
            )
            dibuat += 1

        kata = "akan diimpor" if dry else "diimpor"
        self.stdout.write(
            self.style.SUCCESS(f"Selesai. {dibuat} UMKM {kata}, {dilewati} dilewati.")
        )
